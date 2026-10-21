"""Search and import official open-data exports. Imports are atomic and keep their provenance."""

import asyncio
import csv
import hashlib
import json
import re
import tempfile
import xml.etree.ElementTree as ET
import zipfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from urllib.parse import urlparse

from . import anaf
from .common import fold
from .evidence import now

API = "https://data.gov.ro/api/3/action/"
MAX_DOWNLOAD = 2_000_000_000


def day(value):
    return date.fromisoformat(value).isoformat()


def key(value):
    return re.sub(r"[^a-z0-9]", "", fold(str(value or "")))


async def action(client, name, **params):
    response = await client.get(API + name, params=params)
    response.raise_for_status()
    data = response.json()
    if not data.get("success"):
        raise ValueError("data.gov.ro could not complete the request")
    return data["result"]


async def catalog(client, query, limit=10):
    result = await action(client, "package_search", q=query, rows=limit, sort="metadata_modified desc")
    return {"total": result["count"], "datasets": [
        {"id": d["id"], "title": d["title"], "url": "https://data.gov.ro/dataset/" + d["name"],
         "publisher": (d.get("organization") or {}).get("title"),
         "metadata_modified": d.get("metadata_modified"), "license": d.get("license_title"),
         "resources": [{"id": r["id"], "name": r.get("name"), "format": r.get("format"), "url": r["url"]}
                       for r in d.get("resources", [])]}
        for d in result["results"]],
        "note": "Metadata modification dates do not establish the period covered by a file. Read its title and notes."}


def _rows(path):
    suffix = Path(path).suffix.lower()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        with zipfile.ZipFile(path) as archive:
            if sum(i.file_size for i in archive.infolist()) > 4_000_000_000:
                raise ValueError("The workbook expands beyond the import limit")
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            # SEAP adds a second 'Legenda' sheet. It contains no records.
            sheet = workbook.worksheets[0]
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, ())
            for values in iterator:
                yield dict(zip(headers, values))
        finally:
            workbook.close()
    elif suffix == ".ods":
        ns = {"t": "urn:oasis:names:tc:opendocument:xmlns:table:1.0",
              "x": "urn:oasis:names:tc:opendocument:xmlns:text:1.0"}
        with zipfile.ZipFile(path) as archive:
            if archive.getinfo("content.xml").file_size > 4_000_000_000:
                raise ValueError("The workbook expands beyond the import limit")
            with archive.open("content.xml") as stream:
                headers = None
                for _, element in ET.iterparse(stream, events=("end",)):
                    if element.tag == f"{{{ns['t']}}}table":
                        break  # only the data sheet, not the legend
                    if element.tag != f"{{{ns['t']}}}table-row":
                        continue
                    values = []
                    for cell in element:
                        repeat = min(int(cell.get(f"{{{ns['t']}}}number-columns-repeated", "1")), 512)
                        text = " ".join("".join(p.itertext()) for p in cell.findall("x:p", ns))
                        values.extend([text] * repeat)
                    if headers is None:
                        headers = values
                    elif any(values):
                        yield dict(zip(headers, values))
                    element.clear()
    elif suffix in (".csv", ".txt"):
        with open(path, encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(65536)
            stream.seek(0)
            header = sample.splitlines()[0] if sample else ""
            delimiter = max(("^", ";", "\t", ","), key=header.count)
            yield from csv.DictReader(stream, delimiter=delimiter)
    else:
        raise ValueError("Import a CSV, XLSX or ODS export; legacy XLS must be converted first")


def _get(row, *names):
    for name in names:
        value = row.get(key(name))
        if value is not None and str(value).strip():
            return str(value).strip()
    return ""


def _money(value):
    value = str(value).replace("\xa0", "").replace(" ", "")
    if "," in value:
        value = value.replace(".", "").replace(",", ".")
    number = Decimal(value)
    if not number.is_finite() or number < 0:
        raise ValueError("Invalid contract value")
    return str(number)


def _signed(value):
    if not value:
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(value[:10], fmt).date().isoformat()
        except ValueError:
            pass
    raise ValueError("Unrecognized contract date")


def import_file(store, path, kind, source_url, published_on, period_start=None, period_end=None,
                dataset_id=None, license=None):
    if kind not in ("onrc", "caen", "procurement"):
        raise ValueError("kind must be onrc, caen or procurement")
    published_on = day(published_on)
    period_start = day(period_start) if period_start else None
    period_end = day(period_end) if period_end else None
    if period_start and period_end and period_start > period_end:
        raise ValueError("The period start must precede its end")
    if not source_url.startswith("https://"):
        raise ValueError("Provide the https URL of the source export")
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    dataset_id = dataset_id or hashlib.sha256(source_url.encode()).hexdigest()[:24]
    table = {"onrc": "companies", "caen": "company_caen", "procurement": "contracts"}[kind]
    count = skipped = 0
    with store.connect() as db:
        previous = db.execute("select kind from datasets where id=?", (dataset_id,)).fetchone()
        if previous and previous["kind"] != kind:
            raise ValueError("This dataset ID was previously imported with a different kind")
        db.execute(f"delete from {table} where dataset_id=?", (dataset_id,))
        for raw in _rows(path):
            if not any(v is not None and str(v).strip() for v in raw.values()):
                continue
            row = {key(k): v for k, v in raw.items() if k is not None}
            try:
                if kind == "onrc":
                    cui = anaf.clean_cui(_get(row, "cui", "cif"))
                    name = _get(row, "denumire", "name", "denumire firma")
                    if not name:
                        raise ValueError("Missing company name")
                    county = _get(row, "adr_judet", "judet", "county")
                    db.execute("insert or replace into companies values (?,?,?,?,?,?,?,?,?)", (
                        dataset_id, cui, name, " ".join(re.findall(r"[a-z0-9]+", fold(name))), county,
                        fold(county), _get(row, "caen", "cod caen"),
                        _get(row, "cod_inmatriculare", "nr reg com", "registration_number"),
                        _get(row, "address", "adresa") or " ".join(_get(row, field) for field in
                            ("adr_localitate", "adr_den_strada", "adr_nr_strada")).strip()))
                elif kind == "caen":
                    registration = _get(row, "cod_inmatriculare", "registration_number")
                    caen = _get(row, "cod_caen_autorizat", "cod caen", "caen")
                    if caen:
                        caen = caen.zfill(4)
                    if not registration or not re.fullmatch(r"\d{4}", caen):
                        raise ValueError("Missing registration number or CAEN code")
                    db.execute("insert or replace into company_caen values (?,?,?,?)", (
                        dataset_id, registration, caen, _get(row, "ver_caen_autorizat", "version")))
                else:
                    supplier_cui = anaf.clean_cui(_get(row, "cui ofertant castigator", "cui furnizor",
                                                     "cif ofertant", "supplier_cui"))
                    buyer = _get(row, "autoritate contractanta", "denumire autoritate contractanta", "buyer")
                    title = _get(row, "denumire contract", "denumire achizitie", "denumire cpv", "title")
                    value = _money(_get(row, "valoare contract (ron)", "valoare cumparare directa", "value"))
                    signed = _signed(_get(row, "data contract", "data finalizare", "signed_on"))
                    if not buyer or not title:
                        raise ValueError("Missing buyer or contract title")
                    identity = [_get(row, "numar anunt atribuire", "id"), _get(row, "numar contract"),
                                _get(row, "numar lot"), str(supplier_cui), buyer]
                    if not identity[0] and not identity[1]:
                        identity.extend([signed, title, value])
                    record_id = hashlib.sha256(json.dumps(identity).encode()).hexdigest()
                    db.execute("insert or replace into contracts values (?,?,?,?,?,?,?,?,?)", (
                        dataset_id, record_id, supplier_cui,
                        _get(row, "ofertant castigator", "denumire furnizor", "supplier"), buyer, title, signed,
                        value, _get(row, "currency", "moneda") or "RON"))
                count += 1
            except (ValueError, InvalidOperation):
                skipped += 1
        if not count:
            raise ValueError("No usable rows found. Check the file type, headers and import kind; previous data was kept")
        actual = db.execute(f"select count(*) from {table} where dataset_id=?", (dataset_id,)).fetchone()[0]
        db.execute("insert or replace into datasets values (?,?,?,?,?,?,?,?,?,?,?)", (
            dataset_id, kind, source_url, published_on, period_start, period_end, now(), actual,
            skipped, digest.hexdigest(), license))
    return {"id": dataset_id, "kind": kind, "rows": actual, "skipped": skipped,
            "published_on": published_on, "sha256": digest.hexdigest(),
            "note": "Rows with missing or invalid identifiers were skipped. This is an export, not a live registry query."}


async def sync_resource(client, store, resource_id, kind, published_on, period_start=None, period_end=None):
    day(published_on)
    if period_start:
        day(period_start)
    if period_end:
        day(period_end)
    resource = await action(client, "resource_show", id=resource_id)
    package = await action(client, "package_show", id=resource["package_id"])
    url = resource["url"]
    parsed = urlparse(url)
    if parsed.scheme != "https" or parsed.hostname != "data.gov.ro":
        raise ValueError("This importer downloads files hosted on data.gov.ro only")
    suffix = Path(parsed.path).suffix.lower()
    if suffix not in (".csv", ".txt", ".xlsx", ".ods"):
        raise ValueError("Select a CSV, XLSX or ODS resource")
    with tempfile.TemporaryDirectory(prefix="romania-mcp-") as folder:
        path = Path(folder) / ("export" + suffix)
        async with client.stream("GET", url, timeout=180) as response:
            response.raise_for_status()
            size = 0
            with path.open("wb") as target:
                async for chunk in response.aiter_bytes():
                    size += len(chunk)
                    if size > MAX_DOWNLOAD:
                        raise ValueError("Export exceeds the 2 GB download limit")
                    target.write(chunk)
        return await asyncio.to_thread(import_file, store, path, kind, url, published_on,
                                       period_start, period_end, resource_id, package.get("license_title"))


def manifests(store, kind=None):
    with store.connect() as db:
        query = "select * from datasets" + (" where kind=?" if kind else "")
        return [dict(r) for r in db.execute(query + " order by published_on desc", (kind,) if kind else ())]


def search(store, name, county=None, caen=None, limit=20, offset=0, caen_version=None):
    tokens = re.findall(r"[a-z0-9]+", fold(name))
    if not tokens:
        raise ValueError("Enter at least part of a company name")
    clauses = ["c.name_key like ? escape '\\'" for _ in tokens]
    args = ["%" + t + "%" for t in tokens]
    if county:
        clauses.append("c.county_key=?")
        args.append(fold(county))
    if caen:
        condition = """exists(select 1 from company_caen cc join datasets cd on cd.id=cc.dataset_id
            where cc.registration_number=c.registration_number and cc.caen=? and cd.published_on=c.published_on"""
        if caen_version:
            condition += " and cc.version=?"
        condition += ")"
        if caen_version:
            clauses.append(condition)
            args.extend([caen, caen_version])
        else:
            clauses.append("(c.caen=? or " + condition + ")")
            args.extend([caen, caen])
    with store.connect() as db:
        prefix = """with ranked as (select c.*, d.published_on,
            row_number() over(partition by c.cui order by d.published_on desc,d.imported_at desc) rank
            from companies c join datasets d on d.id=c.dataset_id
            where d.published_on=(select max(published_on) from datasets where kind='onrc'))
            select c.* from ranked c where c.rank=1 and """ + " and ".join(clauses)
        total = db.execute("select count(*) from (" + prefix + ")", args).fetchone()[0]
        rows = [dict(r) for r in db.execute(prefix + " order by c.name,c.cui limit ? offset ?",
                                          (*args, limit, offset))]
    for row in rows:
        row.pop("rank", None)
        row.pop("name_key", None)
        row.pop("county_key", None)
    return {"companies": rows, "total": total, "offset": offset, "caen_version": caen_version,
            "datasets": manifests(store, "onrc"),
            "caen_datasets": manifests(store, "caen") if caen else [],
            "status": "available" if manifests(store, "onrc") else "not_loaded",
            "note": "Choose the matching CUI and verify it with ANAF. Search uses the latest imported ONRC snapshot. "
                    "Authorized CAEN filters require its matching snapshot; specify caen_version when the revision matters."}


def procurement(store, cui, since=None, until=None, limit=20, offset=0):
    clauses, args = ["supplier_cui=?"], [cui]
    if since:
        clauses.append("signed_on>=?")
        args.append(day(since))
    if until:
        clauses.append("signed_on<=?")
        args.append(day(until))
    with store.connect() as db:
        # Repeated exports of the same record should not inflate totals.
        query = """with ranked as (select c.*,d.published_on,
            row_number() over(partition by record_id,supplier_cui order by d.published_on desc,d.imported_at desc) rank
            from contracts c join datasets d on d.id=c.dataset_id)
            select * from ranked where rank=1 and """ + " and ".join(clauses)
        totals, total = {}, 0
        for row in db.execute(query, args):
            currency = row["currency"]
            totals[currency] = totals.get(currency, Decimal(0)) + Decimal(row["value"])
            total += 1
        rows = [dict(r) for r in db.execute(query + " order by signed_on desc,record_id desc limit ? offset ?",
                                          (*args, limit, offset))]
    for row in rows:
        row.pop("rank", None)
    return {"cui": cui, "status": "available" if manifests(store, "procurement") else "not_loaded",
            "total": total, "value_by_currency": {k: str(v) for k, v in totals.items()},
            "contracts": rows, "offset": offset, "datasets": manifests(store, "procurement"),
            "note": "Totals cover imported records only. No matches does not prove the company has no public contracts. "
                    "Joint bids without a single valid supplier CUI are excluded; contract values are not payments."}
